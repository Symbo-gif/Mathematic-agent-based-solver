# Copyright 2025 Damien Davison & Michael Maillet, Recursive AI Devs
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#     http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.

"""
DOCX to Readable Format Converter
=================================
Converts all .docx files in the Reference Documents folder to readable .md (Markdown) files.

Requirements:
    pip install python-docx

Usage:
    python convert_docs.py
"""

import os
import sys
from pathlib import Path

try:
    from docx import Document
    from docx.opc.exceptions import PackageNotFoundError
except ImportError:
    print("ERROR: python-docx is not installed.")
    print("Please run: pip install python-docx")
    sys.exit(1)


def extract_docx_to_markdown(docx_path: Path) -> str:
    """
    Extract text from a .docx file and convert to Markdown format.
    Preserves headings, paragraphs, lists, and tables.
    """
    try:
        doc = Document(str(docx_path))
    except PackageNotFoundError:
        return f"[ERROR: Could not open {docx_path.name}]"
    except Exception as e:
        return f"[ERROR: {e}]"

    markdown_lines = []

    for element in doc.element.body:
        # Handle paragraphs
        if element.tag.endswith('p'):
            para = None
            for p in doc.paragraphs:
                if p._element == element:
                    para = p
                    break

            if para is None:
                continue

            text = para.text.strip()
            if not text:
                markdown_lines.append("")
                continue

            # Check for heading styles
            style_name = para.style.name if para.style else ""

            if "Heading 1" in style_name or "Title" in style_name:
                markdown_lines.append(f"# {text}")
            elif "Heading 2" in style_name:
                markdown_lines.append(f"## {text}")
            elif "Heading 3" in style_name:
                markdown_lines.append(f"### {text}")
            elif "Heading 4" in style_name:
                markdown_lines.append(f"#### {text}")
            elif "List" in style_name or text.startswith(('-', '•', '*', '–')):
                # Clean up list markers
                clean_text = text.lstrip('-•*– ')
                markdown_lines.append(f"- {clean_text}")
            else:
                # Check for bold/emphasis in runs
                formatted_text = ""
                for run in para.runs:
                    run_text = run.text
                    if run.bold and run.italic:
                        formatted_text += f"***{run_text}***"
                    elif run.bold:
                        formatted_text += f"**{run_text}**"
                    elif run.italic:
                        formatted_text += f"*{run_text}*"
                    else:
                        formatted_text += run_text

                markdown_lines.append(formatted_text if formatted_text else text)

        # Handle tables
        elif element.tag.endswith('tbl'):
            for table in doc.tables:
                if table._element == element:
                    markdown_lines.append("")
                    for i, row in enumerate(table.rows):
                        cells = [cell.text.strip().replace('\n', ' ') for cell in row.cells]
                        markdown_lines.append("| " + " | ".join(cells) + " |")
                        if i == 0:
                            # Add header separator
                            markdown_lines.append("| " + " | ".join(["---"] * len(cells)) + " |")
                    markdown_lines.append("")
                    break

    return "\n".join(markdown_lines)


def convert_all_docx(root_folder: Path, output_format: str = "md"):
    """
    Find and convert all .docx files in the given folder and subfolders.
    """
    docx_files = list(root_folder.rglob("*.docx"))

    if not docx_files:
        print("No .docx files found.")
        return

    print(f"Found {len(docx_files)} .docx file(s) to convert.\n")

    converted = 0
    failed = 0

    for docx_path in docx_files:
        # Skip temporary Word files
        if docx_path.name.startswith("~$"):
            continue

        output_path = docx_path.with_suffix(f".{output_format}")

        print(f"Converting: {docx_path.name}")

        try:
            content = extract_docx_to_markdown(docx_path)

            if content.startswith("[ERROR"):
                print(f"  FAILED: {content}")
                failed += 1
                continue

            # Add source header
            header = f"<!-- Converted from: {docx_path.name} -->\n\n"

            with open(output_path, "w", encoding="utf-8") as f:
                f.write(header + content)

            print(f"  -> {output_path.name}")
            converted += 1

        except Exception as e:
            print(f"  FAILED: {e}")
            failed += 1

    print(f"\n{'='*50}")
    print(f"Conversion complete!")
    print(f"  Converted: {converted}")
    print(f"  Failed: {failed}")
    print(f"  Total: {len(docx_files)}")


def main():
    # Get the project root (parent of scripts folder)
    script_dir = Path(__file__).parent
    project_root = script_dir.parent
    reference_folder = project_root / "Reference Documents"

    if not reference_folder.exists():
        print(f"ERROR: Reference Documents folder not found at {reference_folder}")
        print("Make sure this script is in the 'scripts' folder of the project")
        sys.exit(1)

    print("=" * 50)
    print("DOCX to Markdown Converter")
    print("=" * 50)
    print(f"Scanning: {reference_folder}\n")

    convert_all_docx(reference_folder, output_format="md")

    print("\nAll .md files are now readable by Claude Code!")


if __name__ == "__main__":
    main()
